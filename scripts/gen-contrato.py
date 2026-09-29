# -*- coding: utf-8 -*-
"""Gera o contrato Aurum Tax NCM em .docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "docs/CONTRATO - Aurum Tax NCM (Licenca de Uso Desktop Local).docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

sec = doc.sections[0]
sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3); sec.right_margin = Cm(3)

def titulo(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.bold = True; r.font.size = Pt(14)
    return p

def subtitulo(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p

def clausula(texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True; r.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    return p

def par(texto, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    p.add_run(texto)
    return p

def bullet(texto):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(texto)
    return p

titulo("CONTRATO DE LICENÇA DE USO DE SOFTWARE DESKTOP — AURUM TAX NCM")
subtitulo("Licença de uso local (instalação na máquina do cliente) · Sem armazenamento em nuvem · Versão 1.0 — Setembro/2026")

par("Por este instrumento particular, de um lado:")
par("CONTRATADA: AURUM BIT LABS & STUDIOS LTDA, sociedade empresária limitada, inscrita no CNPJ/MF sob o nº 44.366.379/0001-23, com nome fantasia “AURUM BIT STUDIOS”, com sede na Rua Cônego de Castro, nº 5185, Loja 01, bairro Canindezinho, Fortaleza/CE, CEP 60.763-433, neste ato representada por seu sócio-administrador, Sr. Francisco Davi Carneiro Brito, titular e desenvolvedora do software “Aurum Tax NCM”, doravante denominada simplesmente “CONTRATADA” ou “AURUM”;", "CONTRATADA: ")
par("e, de outro lado:")
par("CONTRATANTE: a pessoa física ou jurídica que adquire a licença de uso do software e adere a este Contrato por meio do Termo de Aceite exibido na instalação local, qualificada pelos dados informados no ato da aquisição/ativação (nome/razão social, CPF/CNPJ, e-mail e demais dados de cadastro), os quais constituem parte integrante e indissociável deste instrumento, doravante denominada simplesmente “CONTRATANTE”.", "CONTRATANTE: ")
par("As partes acima identificadas, doravante denominadas em conjunto “Partes” e isoladamente “Parte”, têm entre si justo e contratado o presente Contrato de Licença de Uso de Software Desktop (“Contrato”), que se regerá pelas cláusulas e condições a seguir, mediante as seguintes considerações:")
par("CONSIDERANDO que a CONTRATADA é a titular dos direitos sobre o software Aurum Tax NCM, programa de computador desktop (Windows/Linux/macOS via Electron) de classificação e análise fiscal da Reforma Tributária, que auxilia na classificação de NCM/NBS, definição de CST de IBS/CBS e cClassTrib, simulação de cálculos, importação e análise de SPED Fiscal e de XMLs de NF-e/NFC-e, gestão de produtos e emissão de relatórios;")
par("CONSIDERANDO que o software é instalado e executado 100% (cem por cento) localmente na máquina do CONTRATANTE, sem hospedagem, sem conta em nuvem e sem envio da base de dados da CONTRATANTE aos servidores da CONTRATADA (ver Cláusula Quinta);")
par("CONSIDERANDO que a base tributária de referência do software observa a legislação da Reforma Tributária do Consumo, em especial a Emenda Constitucional nº 132/2023, a Lei Complementar nº 214/2025 (institui IBS, CBS e Imposto Seletivo), o Decreto nº 12.955/2026 (regulamenta a CBS), a Resolução CGIBS nº 6/2026 (regulamenta o IBS) e a Nota Técnica 2025.002 da NF-e (documentos fiscais e CST de IBS/CBS), sem prejuízo de outras normas correlatas indicadas na tela “Legislação” do software;")
par("CONSIDERANDO que a aceitação eletrônica/local deste Contrato é condição indispensável para a instalação e o primeiro uso do software;")
par("resolvem as Partes celebrar o presente Contrato, que será regido pelas cláusulas seguintes.")

clausula("CLÁUSULA PRIMEIRA — DAS DEFINIÇÕES")
par("1.1. Para os fins deste Contrato, os termos abaixo, no singular ou no plural, terão os seguintes significados:")
par("“Software” ou “Aurum Tax NCM”: o programa de computador desktop, seus módulos (Calculadora Tributária, Consulta NCM, Classificação individual e em lote, SPED Fiscal, Notas Fiscais XML, Produtos, Tabelas Auxiliares e Legislação), bases de dados embutidas, documentação e atualizações fornecidas pela CONTRATADA sob a marca Aurum Tax NCM.", "“Software” ou “Aurum Tax NCM”: ")
par("“Instalação Local”: a instalação e execução do Software exclusivamente no(s) computador(es) de propriedade/posse da CONTRATANTE, com processamento e guarda de dados nesse mesmo ambiente, sem dependência de servidor remoto da CONTRATADA.", "“Instalação Local”: ")
par("“Usuário”: cada pessoa física autorizada pela CONTRATANTE a operar o Software na máquina licenciada.", "“Usuário”: ")
par("“Licença”: o direito de uso do Software concedido por este Contrato, nos limites da Proposta Comercial/Pedido (número de máquinas, vigência e versão).", "“Licença”: ")
par("“Proposta Comercial” / “Pedido”: o documento (orçamento, proposta, nota, e-mail de contratação ou tela de compra) que identifica a modalidade adquirida, o preço, a forma de pagamento e o escopo da Licença, integrando este Contrato.", "“Proposta Comercial” / “Pedido”: ")
par("“LGPD”: a Lei nº 13.709/2018 (Lei Geral de Proteção de Dados Pessoais) e sua regulamentação.", "“LGPD”: ")

clausula("CLÁUSULA SEGUNDA — DO OBJETO")
par("2.1. O presente Contrato tem por objeto a concessão, pela CONTRATADA à CONTRATANTE, de licença de uso, não exclusiva, intransferível (salvo hipótese da Cláusula Décima Quinta) e nos limites do Pedido, do software Aurum Tax NCM, em modalidade desktop de instalação local, bem como a prestação dos serviços correlatos de fornecimento de atualizações e suporte técnico, na forma deste Contrato.")
par("2.2. Conforme a versão adquirida, o Software abrange: (i) calculadora de IBS/CBS; (ii) consulta de NCM com classificações da Reforma (CST, cClassTrib, base legal e trechos da legislação); (iii) classificação individual e em lote (CSV/Excel); (iv) importação e análise de SPED Fiscal; (v) importação e análise de XMLs de NF-e/NFC-e; (vi) cadastro e gestão de empresas e produtos; (vii) tabelas auxiliares editáveis (CST, cClassTrib, NCM, CFOP, CST ICMS, CST PIS/COFINS); (viii) módulo de legislação com links oficiais; e (ix) exportação de relatórios em PDF/planilha. O escopo efetivo é o descrito na Proposta/Pedido e na versão instalada.")
par("2.3. A licença ora concedida não implica transferência de propriedade, de código-fonte ou de qualquer direito autoral ou de propriedade intelectual sobre o Software, que permanecem integralmente com a CONTRATADA, protegidos pelas Leis nº 9.609/1998 (Software), nº 9.610/1998 (Direitos Autorais) e nº 9.279/1996 (Propriedade Industrial).")
par("2.4. O Software é ferramenta de apoio. Não substitui o julgamento técnico-profissional da CONTRATANTE, na forma da Cláusula Décima Primeira.")

clausula("CLÁUSULA TERCEIRA — DA ACEITAÇÃO NA INSTALAÇÃO LOCAL (TERMO DE ACEITE)")
par("3.1. A instalação e o primeiro uso do Software pressupõem a leitura integral e a aceitação livre, expressa e inequívoca deste Contrato pela CONTRATANTE, por meio do Termo de Aceite exibido na tela de instalação/primeira execução (caixa de seleção + botão “Li e Aceito”, ou meio equivalente), sem o qual o instalador não prossegue e o Software não é liberado.")
par("3.2. As Partes reconhecem a validade, a eficácia e a força probante da manifestação de vontade realizada por meio eletrônico/local, nos termos do art. 10, §2º, da Medida Provisória nº 2.200-2/2001, da Lei nº 14.063/2020 e dos arts. 104, 107 e 219 do Código Civil, dispensada a assinatura física e a presença de testemunhas.")
par("3.3. O Software registra localmente os elementos comprobatórios do aceite — versão deste Contrato aceita, data e hora, identificação da máquina/instalação e identificação do Usuário que aceitou —, que servirão como prova hábil e suficiente da contratação, sendo desde já reconhecidos pela CONTRATANTE como meio idôneo de comprovação. A CONTRATADA poderá solicitar a reapresentação do aceite a cada atualização relevante do Contrato.")
par("3.4. A pessoa que realiza o aceite declara possuir capacidade e poderes para obrigar a CONTRATANTE, respondendo pessoalmente caso não os detenha.")
par("3.5. O inteiro teor deste Contrato é entregue em formato Word/PDF (docs) e permanece acessível no instalador (arquivo LICENCA.txt) e no menu “Sobre/Ajuda” do Software.")

clausula("CLÁUSULA QUARTA — DA LICENÇA DE USO, INSTALAÇÃO E RESTRIÇÕES")
par("4.1. A licença autoriza a instalação e o uso do Software exclusivamente para os fins lícitos e próprios da atividade da CONTRATANTE, no limite do Pedido (número de máquinas/usuários, vigência e versão). Salvo previsão expressa no Pedido, a licença é vinculada à(s) máquina(s) indicada(s) no ato da ativação.")
par("4.2. São requisitos e deveres técnicos da CONTRATANTE: (i) prover computador compatível (sistema operacional e recursos indicados na documentação), (ii) manter o ambiente atualizado e livre de malware, (iii) realizar os procedimentos de instalação com privilégios adequados e (iv) aplicar as atualizações recomendadas pela CONTRATADA.")
par("4.3. É expressamente vedado à CONTRATANTE e a seus Usuários, direta ou indiretamente:")
for t in ["(a) copiar, reproduzir ou distribuir o Software além do número de instalações licenciadas;",
"(b) modificar, traduzir, adaptar ou criar obras derivadas do Software;",
"(c) realizar engenharia reversa, descompilação, desmontagem ou qualquer tentativa de obtenção do código-fonte ou da lógica do Software, ressalvadas as hipóteses legais expressas;",
"(d) sublicenciar, ceder, locar, emprestar, revender ou de qualquer forma disponibilizar o Software a terceiros não autorizados;",
"(e) burlar mecanismos de ativação, licenciamento ou controle de cópias;",
"(f) utilizar o Software para finalidade ilícita, fraudulenta, que viole direitos de terceiros, a legislação ou a ordem pública;",
"(g) introduzir códigos maliciosos ou comprometer a segurança do ambiente;",
"(h) remover, ocultar ou alterar marcas, avisos de direitos autorais ou sinais de identificação da CONTRATADA."]:
    bullet(t)
par("4.4. O descumprimento desta Cláusula autoriza a suspensão da licença (bloqueio de ativação/atualizações) e a rescisão por justa causa, sem prejuízo das perdas e danos e das medidas judiciais cabíveis.")

clausula("CLÁUSULA QUINTA — DO FUNCIONAMENTO 100% LOCAL E DA AUSÊNCIA DE NUVEM (PRIVACIDADE POR ARQUITETURA)")
par("5.1. O Aurum Tax NCM é um software desktop de processamento e armazenamento locais. Todos os dados operacionais da CONTRATANTE — empresas, produtos, SPED, XMLs de notas, classificações, reclassificações manuais, emitente e tabelas auxiliares — são gravados exclusivamente na máquina da CONTRATANTE (banco local IndexedDB/Dexie no perfil da aplicação, pasta de dados do usuário — userData/xml — e chaves de localStorage da aplicação), e NÃO são enviados, hospedados, replicados ou armazenados em nuvem pela CONTRATADA.", "Destaque — sem nuvem: ")
par("5.2. A CONTRATADA NÃO tem acesso, NÃO coleta, NÃO visualiza e NÃO trata os dados operacionais inseridos pela CONTRATANTE no Software. Não há conta em nuvem, painel administrativo remoto ou telemetria da base de dados. Por consequência, a CONTRATADA não atua como operadora desses dados para fins da LGPD — o tratamento ocorre exclusivamente no ambiente e sob o controle da CONTRATANTE.")
par("5.3. Exceções pontuais e opcionais, que NÃO constituem armazenamento em nuvem da base: (i) consulta cadastral de CNPJ via BrasilAPI, quando a CONTRATANTE clicar em “Buscar/Puxar dados” (tráfego restrito ao CNPJ digitado); (ii) leitura interna de norma oficial (Planalto/CGIBS) pelo módulo Legislação, que baixa o HTML público da norma para exibição; (iii) acesso externo, em nova aba do navegador, ao Portal da Conformidade Fácil (CFF/SEFAZ-RS); e (iv) verificação de atualizações do Software, se habilitada. Em nenhuma dessas hipóteses a base local da CONTRATANTE é transmitida à CONTRATADA.")
par("5.4. Conexão à internet é necessária apenas para as funções opcionais do item 5.3 e para baixar atualizações. Todas as funções essenciais de classificação, cálculo, SPED, XML, produtos e relatórios operam offline após a instalação.")
par("5.5. Como os dados jamais saem da máquina por ação do Software, a confidencialidade, o controle de acesso físico/lógico ao computador, a gestão de Usuários do sistema operacional e a observância da LGPD perante titulares e autoridades competem integralmente à CONTRATANTE.")

clausula("CLÁUSULA SEXTA — DO BACKUP, DA PORTABILIDADE E DA GUARDA DOS DADOS")
par("6.1. Cabe exclusivamente à CONTRATANTE realizar e testar rotinas próprias de backup dos dados locais (função “Backup completo” em Configurações gera um JSON com empresas, produtos, notas e emitente; os XMLs originais permanecem na pasta de dados do usuário). A CONTRATADA não mantém cópia de segurança dos dados da CONTRATANTE, justamente porque não os recebe.")
par("6.2. Formatação, troca de máquina, falha de disco ou desinstalação sem backup implicam perda definitiva dos dados locais, hipótese pela qual a CONTRATADA não se responsabiliza.")
par("6.3. A CONTRATANTE poderá exportar seus dados a qualquer tempo, nos formatos disponibilizados pelo Software (PDF, CSV/XLSX, JSON de backup).")

clausula("CLÁUSULA SÉTIMA — DO PREÇO, DA FORMA DE PAGAMENTO E DO REAJUSTE")
par("7.1. Pela licença e pelos serviços correlatos, a CONTRATANTE pagará o valor indicado na Proposta Comercial/Pedido, na modalidade ali prevista (licença perpétua de versão e/ou assinatura de atualizações/suporte, pagamento único ou recorrente), acrescido de eventuais adicionais contratados.")
par("7.2. Os pagamentos, quando eletrônicos, são processados por instituições/gateways de terceiros. A CONTRATANTE autoriza, quando aplicável, a cobrança recorrente automática até o cancelamento na forma deste Contrato.")
par("7.3. O não pagamento sujeitará a CONTRATANTE a: (i) multa de 2% (dois por cento), juros de mora de 1% (um por cento) ao mês e correção monetária; (ii) após aviso, suspensão de atualizações, suporte e, se aplicável, da ativação; sem prejuízo da cobrança dos valores devidos.")
par("7.4. Os valores poderão ser reajustados anualmente pela variação do IPCA/IBGE ou índice substituto, e poderão ser revistos mediante aviso prévio de 30 (trinta) dias; não concordando, a CONTRATANTE poderá rescindir antes da vigência do novo preço, mantida a versão já licenciada em caráter perpétuo, quando for o caso.")
par("7.5. Tributos incidentes sobre o preço serão suportados na forma da lei. A emissão de documento fiscal observará a legislação aplicável à CONTRATADA.")

clausula("CLÁUSULA OITAVA — DO PRAZO, DAS ATUALIZAÇÕES E DA RESCISÃO")
par("8.1. Este Contrato vigora a partir do aceite na instalação. A Licença observará o modelo do Pedido: (a) licença perpétua da versão adquirida (uso por prazo indeterminado daquela versão, sem direito automático a versões maiores futuras, salvo contratação); e/ou (b) assinatura de atualizações/suporte por ciclos que se renovam automaticamente, salvo manifestação em contrário.")
par("8.2. Atualizações corretivas e evolutivas da base tributária (ex.: ajustes de CST/cClassTrib, NCM, textos legais) e do Software são disponibilizadas pela CONTRATADA por download/instalador. É dever da CONTRATANTE aplicá-las, pois a legislação da Reforma encontra-se em implantação.")
par("8.3. A CONTRATANTE poderá rescindir a assinatura (quando houver) a qualquer tempo, cessando a renovação subsequente, sem devolução de valores de períodos já iniciados.")
par("8.4. A CONTRATADA poderá rescindir: (i) imotivadamente, com aviso de 30 (trinta) dias (efeitos limitados a atualizações/suporte futuros, mantido o uso da versão perpétua já licenciada, se houver); ou (ii) imediatamente e por justa causa, em caso de violação contratual, uso ilícito ou inadimplência não sanada.")
par("8.5. Encerrado o Contrato, cessam atualizações e suporte (quando por assinatura). Os dados locais permanecem na máquina da CONTRATANTE, que poderá exportá-los livremente; nenhuma eliminação remota é realizada pela CONTRATADA, pois ela não detém os dados.")

clausula("CLÁUSULA NONA — DAS OBRIGAÇÕES DA CONTRATADA")
par("9.1. Obriga-se a CONTRATADA a:")
for t in ["(a) disponibilizar o instalador e a documentação do Software conforme a versão contratada;",
"(b) fornecer atualizações da base tributária de referência e do Software com esforços comercialmente razoáveis, sem garantia de periodicidade mínima;",
"(c) prestar suporte técnico nos canais e horários informados, em regime de melhores esforços (best-effort);",
"(d) manter a arquitetura local sem coleta da base de dados da CONTRATANTE, nos termos da Cláusula Quinta;",
"(e) observar a LGPD quanto aos dados de contato/cadastro da CONTRATANTE de que seja controladora (faturamento, ativação e suporte)."]:
    bullet(t)
par("9.2. As obrigações da CONTRATADA são de meio, e não de resultado, ressalvado o disposto expressamente em contrário.")

clausula("CLÁUSULA DÉCIMA — DAS OBRIGAÇÕES E RESPONSABILIDADES DA CONTRATANTE")
par("10.1. Obriga-se a CONTRATANTE a:")
for t in ["(a) utilizar o Software em conformidade com este Contrato e a legislação;",
"(b) fornecer dados verídicos de cadastro/ativação e mantê-los atualizados;",
"(c) responsabilizar-se integral e exclusivamente pelo conteúdo, exatidão, origem e licitude de todos os dados que inserir no Software, inclusive os de terceiros;",
"(d) obter e manter as bases legais, consentimentos e autorizações necessários ao tratamento dos dados de terceiros que inserir no Software, atuando como controladora desses dados em seu próprio ambiente;",
"(e) manter cópias de segurança próprias e testadas (Cláusula Sexta);",
"(f) zelar pela segurança física e lógica da máquina (senhas, antivírus, controle de Usuários do SO), respondendo pelo uso por seus prepostos;",
"(g) utilizar consultas externas opcionais (BrasilAPI, portais oficiais) somente para finalidades lícitas, com base legal adequada;",
"(h) arcar com equipamentos, conexão e ambiente tecnológico necessários ao uso."]:
    bullet(t)
par("10.2. A CONTRATANTE reconhece que a CONTRATADA não exerce controle sobre os dados por ela inseridos e que é a única responsável perante terceiros e autoridades por tais dados e por seu uso.")

clausula("CLÁUSULA DÉCIMA PRIMEIRA — DA ISENÇÃO DE RESPONSABILIDADE PROFISSIONAL, FISCAL E LEGISLATIVA")
par("11.1. O Software é ferramenta de apoio e NÃO substitui o julgamento técnico, a análise e a responsabilidade profissional da CONTRATANTE e de seus Usuários, notadamente em matéria contábil, fiscal, tributária e jurídica.")
par("11.2. Cálculos, classificações de NCM/NBS, sugestões de CST/cClassTrib, alíquotas de referência de IBS/CBS, relatórios e textos legais exibidos têm caráter meramente auxiliar e dependem dos dados informados e da conferência da CONTRATANTE, não constituindo aconselhamento contábil, fiscal ou jurídico, nem garantia de conformidade, deferimento ou resultado perante o Fisco.")
par("11.3. A base de referência observa a EC nº 132/2023, a LC nº 214/2025, o Decreto nº 12.955/2026, a Resolução CGIBS nº 6/2026 e a NT 2025.002, conforme a versão instalada. A legislação da Reforma encontra-se em regulamentação e transição; em caso de divergência, prevalece o texto oficial (Planalto/CGIBS/SEFAZ). A CONTRATANTE é a única responsável pela conferência, revisão e entrega de obrigações, declarações, cálculos e documentos, bem como por eventuais erros, multas e autuações, ainda que os dados tenham transitado pelo Software.")
par("11.4. A CONTRATADA não se responsabiliza por perdas decorrentes de desatualização normativa entre atualizações, de divergências de interpretação ou de alterações legislativas.")

clausula("CLÁUSULA DÉCIMA SEGUNDA — DO SUPORTE E DA GARANTIA LIMITADA")
par("12.1. O suporte é prestado em regime de melhores esforços, pelos canais e horários informados, podendo priorizar atendimentos por criticidade. O Software é fornecido “no estado em que se encontra” (as is), sem garantias implícitas de adequação a finalidade específica.")
par("12.2. A CONTRATADA garante que o Software executará substancialmente conforme a documentação na versão entregue. A garantia limita-se, à escolha da CONTRATADA, à correção, à substituição da versão ou à devolução proporcional do valor da licença, observada a Cláusula Décima Terceira.")
par("12.3. Não estão cobertos: (i) falhas do ambiente da CONTRATANTE (hardware, SO, permissões, antivírus, disco); (ii) perda de dados por ausência de backup; (iii) uso em desacordo com a documentação; e (iv) indisponibilidade de serviços externos opcionais (BrasilAPI, Planalto, CGIBS, CFF).")

clausula("CLÁUSULA DÉCIMA TERCEIRA — DA LIMITAÇÃO DE RESPONSABILIDADE")
par("13.1. Na máxima extensão permitida pela legislação aplicável, a responsabilidade total e agregada da CONTRATADA decorrente ou relacionada a este Contrato fica limitada ao valor efetivamente pago pela CONTRATANTE à CONTRATADA nos 12 (doze) meses anteriores ao fato gerador (ou ao valor da licença perpétua, se menor).")
par("13.2. Em nenhuma hipótese a CONTRATADA responderá por danos indiretos, lucros cessantes, perda de chance, de receita, de dados locais, de clientela ou de oportunidade, danos reputacionais, multas e autuações fiscais, ou por danos sem nexo causal direto e exclusivo com conduta dolosa ou gravemente culposa comprovada da CONTRATADA.")
par("13.3. A CONTRATADA não responde por: (i) atos de terceiros e serviços externos opcionais; (ii) conteúdo e uso dos dados inseridos pela CONTRATANTE; (iii) decisões tomadas com base nas informações do Software; (iv) caso fortuito e força maior; (v) uso indevido ou em desacordo com este Contrato; e (vi) perda de dados locais por ausência de backup ou por falha do equipamento.")

clausula("CLÁUSULA DÉCIMA QUARTA — DA CONFIDENCIALIDADE E DA PROTEÇÃO DE DADOS (LGPD)")
par("14.1. Cada Parte manterá sigilo sobre informações confidenciais a que tiver acesso em razão deste Contrato, usando-as apenas para cumpri-lo. O dever subsiste por 5 (cinco) anos após o término.")
par("14.2. Dados operacionais da CONTRATANTE: como permanecem exclusivamente na máquina da CONTRATANTE e não são transmitidos à CONTRATADA (Cláusula Quinta), a CONTRATANTE figura como única controladora (e, quando for o caso, operadora em seu próprio ambiente) desses dados, competindo-lhe atender a titulares, à ANPD e às demais autoridades, bem como definir bases legais e prazos de guarda/eliminação.")
par("14.3. Dados de cadastro/ativação da CONTRATANTE (nome, CPF/CNPJ, e-mail, dados de licenciamento e suporte): a CONTRATADA atua como controladora, tratando-os para ativação, faturamento, suporte e cumprimento legal, nos termos de sua Política de Privacidade e da LGPD.")
par("14.4. Cada Parte responde pelas sanções e danos a que der causa por infração às normas de proteção de dados, observados os limites da Cláusula Décima Terceira.")
par("14.5. Contato do Encarregado/DPO da CONTRATADA: Francisco Davi Carneiro Brito, e-mail privacy@aurumbit.com.br, telefone (85) 9665-2002.")

clausula("CLÁUSULA DÉCIMA QUINTA — DAS DISPOSIÇÕES GERAIS")
par("15.1. A tolerância quanto a descumprimento não implica novação ou renúncia.")
par("15.2. A nulidade de qualquer cláusula não prejudica as demais.")
par("15.3. A CONTRATANTE não poderá ceder este Contrato sem anuência prévia e escrita da CONTRATADA. A CONTRATADA poderá cedê-lo a empresas de seu grupo ou sucessoras, mediante simples comunicação.")
par("15.4. Este Contrato, a Proposta/Pedido e a documentação da versão constituem o acordo integral quanto ao objeto. Em caso de conflito, prevalece este Contrato.")
par("15.5. As Partes são independentes, sem vínculo societário, trabalhista ou de mandato.")
par("15.6. Comunicações válidas pelos canais cadastrados (e-mail, avisos no instalador/Software, telefone/mensagens), presumindo-se recebidas no envio. É dever da CONTRATANTE manter contatos atualizados.")
par("15.7. A CONTRATADA poderá alterar este Contrato e o Software mediante nova versão e aviso; o uso continuado após atualização com novo aceite implica concordância. Não concordando, a CONTRATANTE poderá manter a versão anterior já licenciada (quando perpétua) ou rescindir a assinatura.")

clausula("CLÁUSULA DÉCIMA SEXTA — DA LEGISLAÇÃO APLICÁVEL E DO FORO")
par("16.1. Este Contrato é regido pelas leis da República Federativa do Brasil, em especial a Lei nº 9.609/1998 (Software), a Lei nº 9.610/1998, o Código Civil, o Marco Civil da Internet (Lei nº 12.965/2014) no que couber, a LGPD e, quanto à matéria tributária de referência, a EC nº 132/2023, a LC nº 214/2025, o Decreto nº 12.955/2026, a Resolução CGIBS nº 6/2026 e a NT 2025.002.")
par("16.2. As Partes elegem o foro da Comarca de Fortaleza/CE, com renúncia a qualquer outro, ressalvada, em relação à CONTRATANTE consumidora pessoa física, a competência que a lei lhe assegurar.")
par("16.3. As Partes poderão, de comum acordo, submeter conflitos à mediação prévia.")

par("E, por estarem justas e contratadas, as Partes manifestam sua concordância integral por meio do aceite eletrônico/local registrado na instalação, que produz todos os efeitos legais.")
par("Fortaleza/CE, na data registrada eletronicamente no ato do aceite.")
par("CONTRATADA: AURUM BIT LABS & STUDIOS LTDA — CNPJ 44.366.379/0001-23, representada por seu sócio-administrador, Francisco Davi Carneiro Brito.")
par("CONTRATANTE: identificada pelos dados da aquisição/ativação, com aceite local registrado (versão do Contrato, data, hora e identificação da instalação).")

clausula("ANEXO — TERMO DE ACEITE (TEXTO EXIBIDO NA INSTALAÇÃO LOCAL)")
par("Ao instalar o Aurum Tax NCM, o Usuário declara que leu e aceita o Contrato de Licença de Uso Desktop. Em resumo:")
for t in ["1. Licença de uso do Aurum Tax NCM, software desktop da Aurum Bit Labs & Studios LTDA, nos limites da versão adquirida;",
"2. O software roda 100% na sua máquina. NÃO armazenamos seus dados em nuvem: empresas, produtos, SPED, XMLs e classificações ficam somente no seu computador (banco local + pasta de dados do usuário). A Aurum não recebe nem tem acesso à sua base;",
"3. Internet é usada apenas para funções opcionais (consulta de CNPJ via BrasilAPI, leitura de normas oficiais Planalto/CGIBS, portal CFF e verificação de atualizações) — sua base nunca é enviada à Aurum;",
"4. Você é responsável por seus dados, backups, conferência fiscal (LC 214/2025, Decreto 12.955/2026, Res. CGIBS 6/2026) e pela segurança da máquina;",
"5. O software é ferramenta de apoio e não substitui o julgamento profissional nem garante resultado perante o Fisco;",
"6. O aceite registra versão do Contrato, data/hora e identificação da instalação como prova da contratação."]:
    bullet(t)
par("Clique em “Li e Aceito” para prosseguir com a instalação, ou “Recusar” para encerrar.")

doc.save(OUT)
print(f"OK -> {OUT}")
